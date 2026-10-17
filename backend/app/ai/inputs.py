"""Turn a stored file into Gemini input parts."""

import io
from pathlib import Path

from google.genai import types
from openpyxl import load_workbook
from pypdf import PdfReader
from pypdf.errors import PdfReadError

PDF = "application/pdf"
IMAGE_TYPES = {".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".png": "image/png", ".webp": "image/webp"}
SHEET_TYPES = {".xlsx", ".xlsm"}


class UnreadableFile(ValueError):
    pass


def page_count(filename: str, data: bytes) -> int:
    if Path(filename).suffix.lower() != ".pdf":
        return 1
    try:
        return len(PdfReader(io.BytesIO(data)).pages)
    except PdfReadError as e:
        raise UnreadableFile(f"{filename} is not a readable PDF") from e


def _sheet_text(data: bytes) -> str:
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    out: list[str] = []
    for i, ws in enumerate(wb.worksheets, start=1):
        out.append(f"=== Sheet {i} (treat as page {i}): {ws.title} ===")
        for row in ws.iter_rows(values_only=True):
            if any(v is not None for v in row):
                out.append("\t".join("" if v is None else str(v) for v in row))
    return "\n".join(out)


def document_parts(filename: str, data: bytes) -> list[types.Part]:
    suffix = Path(filename).suffix.lower()
    header = types.Part.from_text(text=f"File name: {filename}")
    if suffix == ".pdf":
        return [header, types.Part.from_bytes(data=data, mime_type=PDF)]
    if suffix in IMAGE_TYPES:
        return [header, types.Part.from_bytes(data=data, mime_type=IMAGE_TYPES[suffix])]
    if suffix in SHEET_TYPES:
        return [header, types.Part.from_text(text=_sheet_text(data))]
    return [header, types.Part.from_text(text=data.decode("utf-8", errors="replace"))]
